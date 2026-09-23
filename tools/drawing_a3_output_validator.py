#!/usr/bin/env python3
"""TAKY Drawing Engine - ISO A3 absolute output validator."""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any, Dict

MM_PER_INCH = 25.4
PT_PER_INCH = 72.0
EMU_PER_INCH = 914400
A3 = {
    "portrait": (297.0, 420.0),
    "landscape": (420.0, 297.0),
}


def target(orientation: str, dpi: int = 300) -> Dict[str, Any]:
    orientation = orientation.lower()
    if orientation not in A3:
        raise ValueError("ORIENTATION_MUST_BE_PORTRAIT_OR_LANDSCAPE")
    w_mm, h_mm = A3[orientation]
    return {
        "orientation": orientation,
        "width_mm": w_mm,
        "height_mm": h_mm,
        "width_pt": w_mm / MM_PER_INCH * PT_PER_INCH,
        "height_pt": h_mm / MM_PER_INCH * PT_PER_INCH,
        "width_emu": round(w_mm / MM_PER_INCH * EMU_PER_INCH),
        "height_emu": round(h_mm / MM_PER_INCH * EMU_PER_INCH),
        "width_px": round(w_mm / MM_PER_INCH * dpi),
        "height_px": round(h_mm / MM_PER_INCH * dpi),
        "dpi": dpi,
    }


def _result(fmt: str, ok: bool, findings: list[dict], **extra: Any) -> Dict[str, Any]:
    return {"format": fmt, "ok": ok, "findings": findings, **extra}


def validate_pdf(path: str | Path, orientation: str = "landscape", tolerance_pt: float = 0.75) -> Dict[str, Any]:
    import pymupdf as fitz
    t = target(orientation)
    doc = fitz.open(path)
    findings = []
    for index, page in enumerate(doc):
        r = page.rect
        if abs(r.width - t["width_pt"]) > tolerance_pt or abs(r.height - t["height_pt"]) > tolerance_pt:
            findings.append({
                "code": "PDF_PAGE_SIZE_MISMATCH",
                "page": index + 1,
                "actual_pt": [r.width, r.height],
                "expected_pt": [t["width_pt"], t["height_pt"]],
            })
    return _result("pdf", not findings, findings, page_count=doc.page_count)


def validate_png(path: str | Path, orientation: str = "landscape", dpi: int = 300) -> Dict[str, Any]:
    from PIL import Image
    t = target(orientation, dpi)
    img = Image.open(path)
    findings = []
    dx=abs(img.size[0]-t["width_px"])
    dy=abs(img.size[1]-t["height_px"])
    if dx>1 or dy>1:
        findings.append({
            "code": "PNG_PIXEL_SIZE_MISMATCH",
            "actual_px": list(img.size),
            "expected_px": [t["width_px"], t["height_px"]],
            "tolerance_px": 1,
        })
    elif dx or dy:
        findings.append({
            "code": "PNG_PIXEL_ROUNDING_TOLERANCE",
            "severity": "WARNING",
            "actual_px": list(img.size),
            "expected_px": [t["width_px"], t["height_px"]],
            "tolerance_px": 1,
        })
    recorded = img.info.get("dpi")
    if recorded:
        if any(abs(float(v) - dpi) > 1.5 for v in recorded[:2]):
            findings.append({"code":"PNG_DPI_MISMATCH","actual_dpi":list(recorded[:2]),"expected_dpi":dpi})
    else:
        findings.append({"code":"PNG_DPI_METADATA_MISSING","severity":"WARNING"})
    blocking = [f for f in findings if f.get("severity") != "WARNING"]
    return _result("png", not blocking, findings)


def validate_pptx(path: str | Path, orientation: str = "landscape", tolerance_emu: int = 4) -> Dict[str, Any]:
    from pptx import Presentation
    t = target(orientation)
    prs = Presentation(path)
    findings = []
    if abs(prs.slide_width - t["width_emu"]) > tolerance_emu or abs(prs.slide_height - t["height_emu"]) > tolerance_emu:
        findings.append({
            "code":"PPTX_SLIDE_SIZE_MISMATCH",
            "actual_emu":[prs.slide_width, prs.slide_height],
            "expected_emu":[t["width_emu"], t["height_emu"]],
        })
    return _result("pptx", not findings, findings, slide_count=len(prs.slides))


def validate_xlsx(path: str | Path, orientation: str = "landscape") -> Dict[str, Any]:
    from openpyxl import load_workbook
    wb = load_workbook(path)
    findings = []
    checked = 0
    for ws in wb.worksheets:
        checked += 1
        # openpyxl may surface OOXML paperSize as int or string depending on
        # write/read path. ISO A3 is code 8; compare normalized values.
        if str(ws.page_setup.paperSize) != str(ws.PAPERSIZE_A3):
            findings.append({"code":"XLSX_NOT_A3","sheet":ws.title,"actual":ws.page_setup.paperSize})
        if ws.page_setup.orientation != orientation:
            findings.append({"code":"XLSX_ORIENTATION_MISMATCH","sheet":ws.title,"actual":ws.page_setup.orientation,"expected":orientation})
        if ws.page_setup.fitToWidth != 1 or ws.page_setup.fitToHeight != 1:
            findings.append({"code":"XLSX_FIT_NOT_ONE_PAGE","sheet":ws.title,"fit":[ws.page_setup.fitToWidth,ws.page_setup.fitToHeight]})
        try:
            print_area = str(ws.print_area)
        except Exception:
            print_area = ""
        if not print_area:
            findings.append({"code":"XLSX_PRINT_AREA_REQUIRED","sheet":ws.title})
    return _result("xlsx", not findings, findings, sheet_count=checked)


def validate_html(path: str | Path, orientation: str = "landscape") -> Dict[str, Any]:
    text = Path(path).read_text(encoding="utf-8")
    w, h = A3[orientation]
    findings = []
    patterns = {
        "HTML_A3_PAGE_RULE_REQUIRED": rf"@page\s*\{{[^}}]*size\s*:\s*{int(w)}mm\s+{int(h)}mm",
        "HTML_ABSOLUTE_WIDTH_REQUIRED": rf"width\s*:\s*{int(w)}mm",
        "HTML_ABSOLUTE_HEIGHT_REQUIRED": rf"height\s*:\s*{int(h)}mm",
        "HTML_PRINT_COLOR_ADJUST_REQUIRED": r"print-color-adjust\s*:\s*exact",
        "HTML_A3_MARKER_REQUIRED": rf'data-page-size=["\']A3["\'][^>]*data-page-orientation=["\']{orientation}["\']',
    }
    for code, pattern in patterns.items():
        if not re.search(pattern, text, flags=re.I|re.S):
            findings.append({"code":code})
    return _result("html", not findings, findings)


VALIDATORS = {
    ".pdf": validate_pdf,
    ".png": validate_png,
    ".pptx": validate_pptx,
    ".xlsx": validate_xlsx,
    ".html": validate_html,
    ".htm": validate_html,
}


def validate(path: str | Path, orientation: str = "landscape", dpi: int = 300) -> Dict[str, Any]:
    p = Path(path)
    fn = VALIDATORS.get(p.suffix.lower())
    if not fn:
        return _result(p.suffix.lower().lstrip("."), False, [{"code":"UNSUPPORTED_FORMAT"}])
    if p.suffix.lower() == ".png":
        return fn(p, orientation, dpi)
    return fn(p, orientation)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--orientation", default="landscape", choices=["landscape","portrait"])
    parser.add_argument("--dpi", type=int, default=300)
    args = parser.parse_args()
    result = validate(args.path, args.orientation, args.dpi)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["ok"] else 2)


if __name__ == "__main__":
    main()
