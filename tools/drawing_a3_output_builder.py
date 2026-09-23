#!/usr/bin/env python3
"""TAKY Drawing Engine - ISO A3 output builders/configurators."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

MM_PER_INCH = 25.4
EMU_PER_INCH = 914400
A3 = {"portrait": (297.0, 420.0), "landscape": (420.0, 297.0)}
PRODUCTION_CLASSES = {"PREVIEW", "FINAL", "USER_FACING"}

def _assert_internal_primitive(artifact_class: str = "EXPERIMENT"):
    cls = str(artifact_class or "EXPERIMENT").upper()
    if cls in PRODUCTION_CLASSES:
        raise RuntimeError("PRODUCTION_OUTPUT_FORBIDDEN_USE_DRAWING_A3_BUNDLE_EXPORTER_WITH_ADMISSION")
    return cls


def dims(orientation: str = "landscape", dpi: int = 300):
    w_mm, h_mm = A3[orientation]
    return {
        "mm": (w_mm, h_mm),
        "emu": (
            round(w_mm / MM_PER_INCH * EMU_PER_INCH),
            round(h_mm / MM_PER_INCH * EMU_PER_INCH),
        ),
        "px": (
            round(w_mm / MM_PER_INCH * dpi),
            round(h_mm / MM_PER_INCH * dpi),
        ),
    }


def html_shell(body_html: str, orientation: str = "landscape") -> str:
    w, h = A3[orientation]
    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
* {{ box-sizing:border-box; -webkit-print-color-adjust:exact !important; print-color-adjust:exact !important; }}
html, body {{ margin:0; padding:0; background:#e9e9e9; }}
@page {{ size:{int(w)}mm {int(h)}mm; margin:0; }}
.a3-page {{
  width:{int(w)}mm;
  height:{int(h)}mm;
  margin:0 auto;
  overflow:hidden;
  position:relative;
  background:white;
  break-after:page;
  page-break-after:always;
}}
@media print {{
  html, body {{ width:{int(w)}mm; background:white; }}
  .a3-page {{ margin:0; box-shadow:none; }}
}}
</style>
</head>
<body>
<section class="a3-page" data-page-size="A3" data-page-orientation="{orientation}">
{body_html}
</section>
</body>
</html>"""


def new_pdf(path: str | Path, orientation: str = "landscape", artifact_class: str = "EXPERIMENT"):
    _assert_internal_primitive(artifact_class)
    import fitz
    w_mm, h_mm = A3[orientation]
    w_pt = w_mm / MM_PER_INCH * 72.0
    h_pt = h_mm / MM_PER_INCH * 72.0
    doc = fitz.open()
    doc.new_page(width=w_pt, height=h_pt)
    doc.save(path)


def new_png(path: str | Path, orientation: str = "landscape", dpi: int = 300, background="white", artifact_class: str = "EXPERIMENT"):
    _assert_internal_primitive(artifact_class)
    from PIL import Image
    w, h = dims(orientation, dpi)["px"]
    img = Image.new("RGB", (w, h), background)
    img.save(path, dpi=(dpi, dpi))


def new_pptx(path: str | Path, orientation: str = "landscape", artifact_class: str = "EXPERIMENT"):
    _assert_internal_primitive(artifact_class)
    from pptx import Presentation
    prs = Presentation()
    w, h = dims(orientation)["emu"]
    prs.slide_width = w
    prs.slide_height = h
    prs.save(path)


def new_xlsx(path: str | Path, orientation: str = "landscape", print_area: str = "A1:Z60", artifact_class: str = "EXPERIMENT"):
    _assert_internal_primitive(artifact_class)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "A3"
    ws["A1"] = "A3 OUTPUT"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = print_area
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True
    wb.save(path)
