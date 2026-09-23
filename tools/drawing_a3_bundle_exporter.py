#!/usr/bin/env python3
"""TAKY Drawing Engine - canonical A3 SVG to multi-format bundle exporter.

One canonical A3 SVG produces HTML / PDF / PNG / PPTX / XLSX.
Every bundle is validated. A failed bundle is regenerated from the same
canonical SVG up to max_attempts. No output file becomes source authority.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
from pathlib import Path
from typing import Dict

import fitz
from PIL import Image
from pptx import Presentation
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

MM_PER_INCH = 25.4
EMU_PER_INCH = 914400
A3 = {"portrait": (297.0, 420.0), "landscape": (420.0, 297.0)}

HERE = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("a3_validator", HERE / "drawing_a3_output_validator.py")
validator = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(validator)


def _dims(orientation: str, dpi: int):
    w_mm, h_mm = A3[orientation]
    return {
        "w_mm": w_mm,
        "h_mm": h_mm,
        "w_pt": w_mm / MM_PER_INCH * 72.0,
        "h_pt": h_mm / MM_PER_INCH * 72.0,
        "w_emu": round(w_mm / MM_PER_INCH * EMU_PER_INCH),
        "h_emu": round(h_mm / MM_PER_INCH * EMU_PER_INCH),
        "w_px": round(w_mm / MM_PER_INCH * dpi),
        "h_px": round(h_mm / MM_PER_INCH * dpi),
    }


def _validate_svg_root(svg_text: str, orientation: str) -> None:
    w_mm, h_mm = A3[orientation]
    normalized = svg_text.replace(" ", "").lower()
    if f'width="{int(w_mm)}mm"'.lower() not in normalized:
        raise ValueError("SVG_A3_WIDTH_REQUIRED")
    if f'height="{int(h_mm)}mm"'.lower() not in normalized:
        raise ValueError("SVG_A3_HEIGHT_REQUIRED")


def _html(svg_text: str, orientation: str) -> str:
    w_mm, h_mm = A3[orientation]
    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<style>
* {{ box-sizing:border-box; -webkit-print-color-adjust:exact !important; print-color-adjust:exact !important; }}
html,body {{ margin:0; padding:0; background:#ddd; }}
@page {{ size:{int(w_mm)}mm {int(h_mm)}mm; margin:0; }}
.a3-page {{
  width:{int(w_mm)}mm;
  height:{int(h_mm)}mm;
  overflow:hidden;
  margin:0 auto;
  background:#fff;
  break-after:page;
  page-break-after:always;
}}
.a3-page > svg {{ display:block; width:100%; height:100%; }}
@media print {{
  html,body {{ width:{int(w_mm)}mm; height:{int(h_mm)}mm; background:#fff; }}
  .a3-page {{ margin:0; }}
}}
</style>
</head>
<body>
<section class="a3-page" data-page-size="A3" data-page-orientation="{orientation}">
{svg_text}
</section>
</body>
</html>"""


def _pdf_from_svg(svg_text: str, out_path: Path, orientation: str) -> None:
    d = _dims(orientation, 300)
    source = fitz.open(stream=svg_text.encode("utf-8"), filetype="svg")
    source_pdf = fitz.open(stream=source.convert_to_pdf(), filetype="pdf")
    out = fitz.open()
    page = out.new_page(width=d["w_pt"], height=d["h_pt"])
    page.show_pdf_page(page.rect, source_pdf, 0, keep_proportion=False)
    out.save(out_path)
    out.close()


def _png_from_pdf(pdf_path: Path, out_path: Path, dpi: int) -> None:
    doc = fitz.open(pdf_path)
    pix = doc[0].get_pixmap(dpi=dpi, alpha=False)
    pix.save(out_path)


def _pptx_from_png(png_path: Path, out_path: Path, orientation: str) -> None:
    d = _dims(orientation, 300)
    prs = Presentation()
    prs.slide_width = d["w_emu"]
    prs.slide_height = d["h_emu"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(png_path),
        0, 0,
        width=prs.slide_width,
        height=prs.slide_height,
    )
    prs.save(out_path)


def _xlsx_from_png(png_path: Path, out_path: Path, orientation: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "A3"
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.orientation = orientation
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True
    ws.print_area = "A1:Z60"
    image = XLImage(str(png_path))
    # Display size is only a worksheet editing surface. Print geometry is
    # controlled by A3 page setup + one-page fit.
    image.width = 1400 if orientation == "landscape" else 990
    image.height = 990 if orientation == "landscape" else 1400
    ws.add_image(image, "A1")
    wb.save(out_path)


def export_once(svg_path: str | Path, out_dir: str | Path, *, orientation: str = "landscape", dpi: int = 300) -> Dict[str, str]:
    svg_path = Path(svg_path)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    svg_text = svg_path.read_text(encoding="utf-8")
    _validate_svg_root(svg_text, orientation)

    paths = {
        "html": out_dir / "board.html",
        "pdf": out_dir / "board.pdf",
        "png": out_dir / "board.png",
        "pptx": out_dir / "board.pptx",
        "xlsx": out_dir / "board.xlsx",
    }
    paths["html"].write_text(_html(svg_text, orientation), encoding="utf-8")
    _pdf_from_svg(svg_text, paths["pdf"], orientation)
    _png_from_pdf(paths["pdf"], paths["png"], dpi)
    _pptx_from_png(paths["png"], paths["pptx"], orientation)
    _xlsx_from_png(paths["png"], paths["xlsx"], orientation)

    return {key: str(value) for key, value in paths.items()}


def validate_bundle(paths: Dict[str, str], orientation: str, dpi: int):
    results = {key: validator.validate(path, orientation, dpi) for key, path in paths.items()}
    return {
        "ok": all(value["ok"] for value in results.values()),
        "results": results,
    }


def export_verified(
    svg_path: str | Path,
    out_dir: str | Path,
    *,
    orientation: str = "landscape",
    dpi: int = 300,
    max_attempts: int = 2,
):
    history = []
    out_dir = Path(out_dir)
    for attempt in range(1, max_attempts + 1):
        if out_dir.exists():
            for child in out_dir.iterdir():
                if child.is_file():
                    child.unlink()
                elif child.is_dir():
                    shutil.rmtree(child)
        paths = export_once(svg_path, out_dir, orientation=orientation, dpi=dpi)
        check = validate_bundle(paths, orientation, dpi)
        history.append({"attempt": attempt, "validation": check})
        if check["ok"]:
            manifest = {
                "status": "PASS",
                "attempts": attempt,
                "orientation": orientation,
                "dpi": dpi,
                "canonical_input": str(Path(svg_path)),
                "outputs": paths,
                "validation": check,
                "history": history,
            }
            (out_dir / "bundle_manifest.json").write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            return manifest

    manifest = {
        "status": "BLOCKED_AFTER_REEXPORT",
        "attempts": max_attempts,
        "orientation": orientation,
        "dpi": dpi,
        "canonical_input": str(Path(svg_path)),
        "history": history,
    }
    (out_dir / "bundle_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return manifest


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("svg")
    parser.add_argument("out_dir")
    parser.add_argument("--orientation", choices=["landscape", "portrait"], default="landscape")
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--max-attempts", type=int, default=2)
    args = parser.parse_args()
    result = export_verified(
        args.svg,
        args.out_dir,
        orientation=args.orientation,
        dpi=args.dpi,
        max_attempts=args.max_attempts,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["status"] == "PASS" else 2)


if __name__ == "__main__":
    main()
